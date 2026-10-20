from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "120_token_option" / "Final_40_Slide_Deck" / "Aurora_Stays_RAG_Final_40_Slides.pptx"
OUTPUT = ROOT / "120_token_option" / "Aurora_Stays_RAG_120_Token_Option_Updated.pptx"


REPLACEMENTS = {
    "TRANSFORMERS_OFFLINE=1 and HF_HUB_OFFLINE=1 prevent runtime model downloads.":
        "TRANSFORMERS_OFFLINE=1 and HF_HUB_OFFLINE=1 prevent runtime model downloads; EXPOSE 7860 receives Codespaces traffic.",
    "Downloads HarishMaths/Hotel-Policy-Embedding at build time and sets HF_HOME=/app/huggingface.":
        "Downloads the Hugging Face embedding model at build time and stores it under HF_HOME=/app/huggingface for reuse.",
    "TRANSFORMERS_OFFLINE=1 and HF_HUB_OFFLINE=1 prevent runtime model downloads; EXPOSE 7860 receives Codespaces traffic.":
        "The Transformers runtime, used through sentence-transformers, loads the cached embedding model; offline flags prevent runtime downloads, and EXPOSE 7860 receives Codespaces traffic.",
    "Docker command: gunicorn -w 1 --timeout 120 -b 0.0.0.0:7860 app:rag_api.":
        "The backend API is exposed on port 7860 for Codespaces traffic.",
    "Binds Streamlit to 0.0.0.0:8501 for Codespaces or container access.":
        "Binds Streamlit to 0.0.0.0:8501 for Codespaces or container access; EXPOSE 8501 makes the service discoverable.",
    "The backend address is supplied through BACKEND_URL; local default is http://localhost:5000.":
        "The frontend does not handle API keys; it sends requests through BACKEND_URL. In Codespaces, this points to the backend forwarded URL on port 7860; local non-Docker tests may use port 5000.",
    "Validated locally with Streamlit on port 8501 and the Flask API on port 5000.":
        "Validated in Codespaces with Streamlit on port 8501 and the Dockerized Flask API on port 7860.",
    "Set BACKEND_URL at runtime so the frontend can call the backend service or forwarded port.":
        "The image defaults BACKEND_URL to http://host.docker.internal:7860; override it with the backend Codespace URL when needed.",
    "Repository: github.com/raovutukoor1/aurora-stay-qna.":
        "Repository: github.com/raovutukoor1/aurora-stay-qna-120-token-option.",
    "Current checked-out commit: b11ab14 — Improve RAG retrieval for breakfast and advance purchase cases.":
        "Current deployment commit: d7a7e3f — fixed the 120-token Docker runtime configuration.",
    "Validated local endpoints: frontend http://localhost:8501 and backend http://127.0.0.1:5000.":
        "Validated public endpoints: frontend port 8501 and backend port 7860; both returned HTTP 200.",
    "Live smoke tests: Streamlit frontend HTTP 200 and backend readiness HTTP 200.":
        "Live Codespaces smoke tests: frontend HTTP 200, backend readiness HTTP 200, answer HTTP 200, and answer-with-evidence HTTP 200 with 3 evidence chunks.",
    "https://didactic-fiesta-4qjv9j6r66pj3jxq6-8501.app.github.dev":
        "https://turbo-zebra-4qqj6wr7pv573xpr-8501.app.github.dev",
    "https://didactic-fiesta-4qjv9j6r66pj3jxq6-5000.app.github.dev":
        "https://turbo-zebra-4qqj6wr7pv573xpr-7860.app.github.dev",
    "Captured: 15 October 2026": "Captured: 20 October 2026",
    "Validated public Codespaces deployment — 15 October 2026.":
        "Validated public Codespaces deployment — 20 October 2026.",
}


def replace_runs(presentation):
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    for old, new in REPLACEMENTS.items():
                        run.text = run.text.replace(old, new)


def main():
    presentation = Presentation(str(INPUT))
    replace_runs(presentation)
    presentation.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
